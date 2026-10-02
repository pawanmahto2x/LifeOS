import os
from docx import Document
from docx.shared import Pt
import re

def create_docx(md_path, docx_path):
    doc = Document()
    
    # Read markdown file
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
            
        # Very basic markdown parsing for headings and bold
        if line.startswith('# '):
            p = doc.add_heading(line[2:].replace('**', ''), level=1)
        elif line.startswith('## '):
            p = doc.add_heading(line[3:].replace('**', ''), level=2)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:].replace('**', ''), level=3)
        else:
            # Add normal paragraph
            # We'll just remove markdown bold stars for simplicity in the output
            clean_text = line.replace('**', '')
            doc.add_paragraph(clean_text)
            
    doc.save(docx_path)
    print(f"Successfully generated {docx_path}")

md_file = "LifeOS_Black_Book_Revised.md"
docx_file = "LifeOS_Black_Book_Final.docx"
create_docx(md_file, docx_file)
