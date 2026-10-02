from docx import Document

doc = Document('LifeOS_Black_Book_Updated.docx')

old_table = doc.tables[6]

# Let's just create a new table at the end of the document, 
# then move its XML element right before the old table.
new_table = doc.add_table(rows=0, cols=5)
new_table.style = 'Table Grid'

data = [
    ['Activity', 'June', 'July', 'August', 'September'],
    ['Requirement Gathering', '████', '', '', ''],
    ['System Analysis & Design', '', '████', '', ''],
    ['Backend API Development', '', '██', '██', ''],
    ['Frontend UI Development', '', '', '████', ''],
    ['Integration & Testing', '', '', '', '███'],
    ['Documentation', '', '', '', '███']
]

for row_data in data:
    row = new_table.add_row()
    for idx, cell_text in enumerate(row_data):
        row.cells[idx].text = cell_text

# Move new_table XML before old_table XML
p = old_table._element
p.addprevious(new_table._element)

# Remove old_table
p.getparent().remove(p)

doc.save('LifeOS_Black_Book_Updated.docx')
print('Gantt table replaced successfully.')
